"""Promotion edge cases that must not lose source evidence or identities."""
import json
from pathlib import Path
import sys

import pytest

pytest.importorskip("ifcopenshell")
pytest.importorskip("openpyxl")

sys.path[:0] = [str(Path(__file__).resolve().parents[1] / "tools")]
from usdaeco_cctv import register_plugins, iter_cameras, sensors_of, camera_type_of
from usdaeco_cctv.importer import Fact, Row, Values, _is_sub, _optics, _table_presets, import_cctv, is_camera
from fixtures import build_baseline, write_cobie, convert
from pxr import Sdf, Usd


@pytest.fixture(scope="module")
def source(tmp_path_factory):
    register_plugins()
    path = tmp_path_factory.mktemp("camera-inputs")
    build_baseline(path / "cameras.ifc")
    convert(path / "cameras.ifc", path / "core.usda")
    return path


def overlay(source, tmp_path, mutate):
    layer = Sdf.Layer.CreateNew(str(tmp_path / "input.usda"))
    layer.subLayerPaths = [str(source / "core.usda")]
    stage = Usd.Stage.Open(layer)
    mutate(stage)
    layer.Save()
    return tmp_path / "input.usda"


@pytest.mark.parametrize("name", ["FOV Pan", "FOV_Camera_Rotation", "FOV 1 Pan"])
def test_proxy_needs_optical_evidence(name):
    assert not is_camera("IfcBuildingElementProxy", [Fact("Identity", "Model", "Camera")])
    assert is_camera("IfcBuildingElementProxy", [Fact("Any", name, 0.)])
    assert not is_camera("IfcWall", [Fact("Any", name, 0.)])


def test_partial_optics_not_blocked():
    values = Values([Fact("Any", "FOV Focal Length Minimum", 3.),
                     Fact("Any", "VideoResolutionWidth", 1920)])
    assert _optics(values) == {}
    assert not values.used


def test_profile_preserves_counters_and_layer_bytes(source, tmp_path):
    import math
    from usdaeco_cctv.importer import STATS
    plain, measured = tmp_path / "plain.usda", tmp_path / "measured.usda"
    expected = import_cctv(source / "core.usda", source / "cameras.ifc", plain)
    timings = {}
    actual = import_cctv(source / "core.usda", source / "cameras.ifc", measured, timings=timings)
    assert actual == expected and set(actual) == set(STATS)
    assert plain.read_bytes() == measured.read_bytes()
    assert set(timings) == {"setup", "sourceRead", "ifcOpen", "ifcRows", "psetRead",
                            "stageOpen", "primTraversal", "quarantineRead", "typeMatching",
                            "cameraAuthoring", "helperMatching", "writing", "total"}
    assert all(math.isfinite(v) and v >= 0 for v in timings.values())
    assert sum(v for k, v in timings.items() if k != "total") <= timings["total"]


def test_cobie_standard_fields_without_pset_category(source, tmp_path):
    from openpyxl import load_workbook
    from usdaeco_cctv.importer import read_cobie, _pose
    workbook = tmp_path / "no-category.xlsx"
    write_cobie(source / "cameras.ifc", workbook)
    book = load_workbook(workbook)
    book["Attribute"].delete_cols(6)
    book.save(workbook)
    book.close()
    fixed = next(r for r in read_cobie(workbook) if r.name == "Fixed")
    pose = _pose(Values(fixed.facts))
    assert pose["pan"] == 15. and pose["tilt"] == 30. and pose["focalLength"] == 4.


def test_duplicate_cobie_component_names_fail_closed(source, tmp_path):
    from openpyxl import load_workbook
    workbook = tmp_path / "duplicate-names.xlsx"
    write_cobie(source / "cameras.ifc", workbook)
    book = load_workbook(workbook)
    book["Component"].cell(3, 1).value = book["Component"].cell(2, 1).value
    book.save(workbook)
    book.close()
    with pytest.raises(ValueError, match="Component.Name must be unique"):
        import_cctv(source / "core.usda", workbook, tmp_path / "kind.usda")
    assert not (tmp_path / "kind.usda").exists()


def test_camera_symbol_choice_is_not_a_subinstance():
    row = Row("", "Device")
    assert not _is_sub(row, [Fact("Visibility", "2D Symbol", "AXIS 2D Symbol : Dome")])
    row.type_name = "AXIS 2D Symbol"
    assert _is_sub(row, [])


def test_internal_formula_is_not_a_focal_driver():
    from usdaeco_cctv.importer import _pose
    values = Values([Fact("Constraints", "_FOV Desired Focal Length", 10000.),
                     Fact("Graphics", "FOV Desired Focal Length", 9.),
                     Fact("Graphics", "FOV Actual Focal Length", 9.)])
    assert _pose(values)["focalLength"] == 9.


def test_unknown_scenario_survives(source, tmp_path):
    def change(stage):
        prim = next(p for p in stage.Traverse() if p.GetName() == "Fixed")
        prim.GetAttribute("aeco:props:Pset_CameraProject:Scenario").Set("CUSTOM_PURPOSE")
    base = overlay(source, tmp_path, change)
    # Use COBie with the same unknown value, otherwise the supplied IFC is
    # the first source of truth and would correctly win the quarantine edit.
    from openpyxl import load_workbook
    write_cobie(source / "cameras.ifc", tmp_path / "cobie.xlsx")
    book = load_workbook(tmp_path / "cobie.xlsx")
    for row in book["Attribute"].iter_rows(min_row=2):
        if row[0].value == "Scenario":
            row[3].value = "CUSTOM_PURPOSE"
    book.save(tmp_path / "cobie.xlsx")
    book.close()
    import_cctv(base, tmp_path / "cobie.xlsx", tmp_path / "kind.usda")
    stage = Usd.Stage.Open(str(tmp_path / "kind.usda"))
    fixed = next(c for c in iter_cameras(stage) if c.GetName() == "Fixed")
    assert fixed.GetAttribute("aeco:cctv:scenario").Get() == ""
    assert fixed.GetAttribute("aeco:props:Pset_CameraProject:Scenario").Get() == "CUSTOM_PURPOSE"


def test_output_refused_and_inputs_not_overwritten(source, tmp_path):
    out = tmp_path / "out.usda"
    out.write_text("keep")
    with pytest.raises(ValueError, match="new file"):
        import_cctv(source / "core.usda", source / "cameras.ifc", out)
    assert out.read_text() == "keep"
    with pytest.raises(ValueError, match="new file"):
        import_cctv(source / "core.usda", source / "cameras.ifc", source / "core.usda")


def test_duplicate_core_identity_refused(source, tmp_path):
    def change(stage):
        prims = [p for p in stage.Traverse() if p.HasAPI("AecoElementAPI")]
        prims[1].GetAttribute("aeco:id").Set(prims[0].GetAttribute("aeco:id").Get())
    base = overlay(source, tmp_path, change)
    with pytest.raises(ValueError, match="duplicate"):
        import_cctv(base, source / "cameras.ifc", tmp_path / "kind.usda")
    assert not (tmp_path / "kind.usda").exists()


def test_absent_catalog_created(source, tmp_path):
    def change(stage):
        prim = next(p for p in stage.Traverse() if p.GetName() == "Fixed")
        prim.GetInherits().SetInherits([])
    base = overlay(source, tmp_path, change)
    stats = import_cctv(base, source / "cameras.ifc", tmp_path / "kind.usda")
    stage = Usd.Stage.Open(str(tmp_path / "kind.usda"))
    fixed = next(c for c in iter_cameras(stage) if c.GetName() == "Fixed")
    cat = camera_type_of(fixed)
    assert cat.GetName().startswith("Camera_") and cat.HasAPI("AecoTypeAPI")
    assert len(sensors_of(fixed)) == 1 and stats["types"] == 3


def test_missing_camera_counted(source, tmp_path):
    def change(stage):
        next(p for p in stage.Traverse() if p.GetName() == "PTZ").SetActive(False)
    base = overlay(source, tmp_path, change)
    stats = import_cctv(base, source / "cameras.ifc", tmp_path / "kind.usda")
    assert stats["cameras"] == 2 and stats["unmatched"] == 1


def test_picture_ownership_by_same_level_proximity(source, tmp_path):
    import ifcopenshell
    model = ifcopenshell.open(str(source / "cameras.ifc"))
    for pset in model.by_type("IfcPropertySet"):
        pset.HasProperties = [p for p in pset.HasProperties if p.Name != "SuperComponent"]
    model.write(str(tmp_path / "proximity.ifc"))
    convert(tmp_path / "proximity.ifc", tmp_path / "core.usda")
    stats = import_cctv(tmp_path / "core.usda", tmp_path / "proximity.ifc", tmp_path / "kind.usda")
    assert stats["subInstancesFolded"] == 2 and stats["unmatched"] == 0


def test_preset_table_sign_and_zoom():
    rows = [{"name": "Preset 1", "value": "12,-25,8"},
            {"name": "Preset_2", "value": {"pan": 45, "tilt": -10, "focalLength": 12, "dwell": 6}}]
    result = _table_presets(json.dumps(rows))
    assert result["Preset_1"] == dict(pan=12., tilt=25., focalLength=8.)
    assert result["Preset_2"]["dwell"] == 6.
    with pytest.raises(ValueError, match="require"):
        _table_presets([{"name": "Broken", "pan": 1}])


def test_ifc_table_reaches_sensor(source, tmp_path):
    import ifcopenshell
    import ifcopenshell.util.element as element
    from ifcopenshell.api import run
    model = ifcopenshell.open(str(source / "cameras.ifc"))
    camera = next(c for c in model.by_type("IfcAudioVisualAppliance") if c.Name == "PTZ")
    pset = run("pset.add_pset", model, product=camera, name="Pset_AudioVisualApplianceTypeCamera")
    pset.HasProperties = [model.create_entity("IfcPropertyTableValue", Name="PanTiltZoomPreset",
        DefiningValues=[model.create_entity("IfcLabel", "Preset_1")],
        DefinedValues=[model.create_entity("IfcText", "10,-20,9")])]
    model.write(str(tmp_path / "table.ifc"))
    import_cctv(source / "core.usda", tmp_path / "table.ifc", tmp_path / "kind.usda")
    stage = Usd.Stage.Open(str(tmp_path / "kind.usda"))
    ptz = next(c for c in iter_cameras(stage) if c.GetName() == "PTZ")
    head = sensors_of(ptz)[0]
    assert head.GetAttribute("aeco:cctvPreset:Preset_1:tilt").Get() == 20.
    assert head.GetAttribute("aeco:cctvPreset:Preset_1:focalLength").Get() == 9.


def test_registered_sync_retains_primary_and_additional_handles(source, tmp_path):
    import os
    import subprocess
    from usdaeco_cctv import ROOT, core_root
    sync = ROOT.parent / "usdaeco-sync/plugins/usdAecoSync/resources"
    if not (sync / "plugInfo.json").exists():
        pytest.skip("optional sync plugin is absent")
    env = dict(os.environ, PXR_PLUGINPATH_NAME=os.pathsep.join([
        str(core_root() / "out/plugins/usdAeco/resources"), str(sync)]))
    env.pop("PYTHONPATH", None)
    code = """
import sys
sys.path.insert(0, sys.argv[1])
from usdaeco_cctv.importer import import_cctv, read_ifc
from usdaeco_cctv import iter_cameras, sensors_of
from pxr import Usd
rows = read_ifc(sys.argv[3])
owner = next(r.uid for r in rows if r.name == 'Fixed')
children = [r.uid for r in rows if r.name.startswith(('FOV-AXIS', 'AXIS 2D Symbol'))]
stats = import_cctv(sys.argv[2], sys.argv[3], sys.argv[4], subinstances={owner: children})
stage = Usd.Stage.Open(sys.argv[4])
fixed = next(p for p in iter_cameras(stage) if p.GetName() == 'Fixed')
sensor = sensors_of(fixed)[0]
assert sensor.HasAPI('AecoHostBindingAPI', 'revit')
refs = [sensor.GetAttribute('aeco:host:revit:ref').Get()] + list(sensor.GetAttribute('aeco:props:cctv:subInstances').Get())
assert {r.rsplit(':', 1)[-1] for r in refs} == set(children)
assert stats['subInstancesFolded'] == 2 and stats['unmatched'] == 0
"""
    result = subprocess.run([sys.executable, "-c", code, str(ROOT / "tools"), str(source / "core.usda"),
                             str(source / "cameras.ifc"), str(tmp_path / "kind.usda")],
                            env=env, text=True, capture_output=True)
    assert result.returncode == 0, result.stderr
