"""Promote IFC, COBie or published USD camera facts into an additive kind layer.

The core stage supplies identity and placement; source properties supply
optics and pose. Only promoted quarantine values are blocked. Neither input
is modified, and an existing output is refused.
"""
import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from itertools import product
import json
import math
import os
from pathlib import Path
import re
import time
import uuid

import numpy as np
from pxr import Gf, Sdf, Usd, UsdGeom

from . import contract, exchange
from .contract import (Fact, Values, _norm, _present, _number, _bool,
                       _ifc_psets, _ifc_facts, pose as _pose, preset_rows as _table_presets)
from .output import publish
from . import __version__, register_plugins, registry, sensors_of

CAMERA_PSET = "Pset_AudioVisualApplianceTypeCamera"
SENSOR = "aeco:cctvSensor:"
STATS = ("cameras", "sensors", "presets", "types", "propsBlocked",
         "subInstancesFolded", "unmatched")
MOUNTS = {17160: "wall", 17161: "ceiling", 17162: "pole", 17165: "corner"}


class _Timings:
    """Exclusive wall-clock stages; never serialized into the kind layer."""
    def __init__(self):
        self.started = self.previous = time.perf_counter()
        self.seconds = {}

    def mark(self, name):
        now = time.perf_counter()
        self.seconds[name] = now - self.previous
        self.previous = now

    def finish(self, destination):
        self.seconds["total"] = time.perf_counter() - self.started
        if destination is not None:
            destination.update(self.seconds)




def identity(value):
    """Canonical UUID for a COBie UUID or compressed IFC GlobalId."""
    value = str(value or "").strip()
    try:
        return str(uuid.UUID(value))
    except ValueError:
        if len(value) == 22:
            import ifcopenshell.guid
            try:
                return str(uuid.UUID(hex=ifcopenshell.guid.expand(value)))
            except (ValueError, IndexError):
                pass
    return value






@dataclass
class Row:
    uid: str
    name: str = ""
    code: str = ""
    type_id: str = ""
    facts: list = field(default_factory=list)
    type_facts: list = field(default_factory=list)
    type_name: str = ""
    ambiguous: bool = False
    parents: tuple = ()  # IFC element ownership, independent of USD containment


def _quarantine(prim):
    if not prim:
        return []
    return [Fact(*a.GetName().split(":", 3)[2:], a.Get())
            for a in prim.GetPropertiesInNamespace("aeco:props")
            if isinstance(a, Usd.Attribute) and a.GetName().startswith("aeco:props:") and a.HasValue()]


def _catalog(prim):
    for path in prim.GetInherits().GetAllDirectInherits():
        candidate = prim.GetStage().GetPrimAtPath(path)
        if candidate and candidate.IsAbstract() and candidate.HasAPI("AecoTypeAPI"):
            return candidate
    return None


def _set(prim, name, value):
    # Sparse occurrence overrides; inherited type optics remain inherited.
    attr = prim.GetAttribute(name)
    if value is not None and (prim.IsAbstract() or not attr.HasAuthoredValueOpinion() or attr.Get() != value):
        attr.Set(exchange.coerce(attr, value))


def _block(prim, used, stats):
    for attr in prim.GetAttributes():
        parts = attr.GetName().split(":", 3)
        if len(parts) == 4 and parts[:2] == ["aeco", "props"]:
            if (_norm(parts[2]), _norm(parts[3])) in used and attr.HasValue():
                attr.Block()
                stats["propsBlocked"] += 1




def read_ifc(path, *, timings=None):
    clock = _Timings()
    import ifcopenshell
    import ifcopenshell.util.element as element
    import ifcopenshell.util.unit as unit
    model = ifcopenshell.open(str(path))
    clock.mark("ifcOpen")
    unit.cache_units(model)
    types, rows = {}, []
    pset_seconds = 0.
    fact_cache = {}

    def facts(entity):
        nonlocal pset_seconds
        started = time.perf_counter()
        result = _ifc_facts(model, entity, cache=fact_cache)
        pset_seconds += time.perf_counter() - started
        return result

    for entity in model.by_type("IfcElement"):
        typ = element.get_type(entity)
        tid = identity(typ.GlobalId) if typ else ""
        if tid and tid not in types:
            types[tid] = facts(typ)
        code = entity.is_a()
        predefined = getattr(entity, "PredefinedType", None) or (getattr(typ, "PredefinedType", None) if typ else None)
        if predefined and predefined not in ("NOTDEFINED", "USERDEFINED"):
            code += "." + predefined
        parents = tuple(sorted({identity(rel.RelatingObject.GlobalId)
                                for rel in (*entity.Decomposes, *entity.Nests)
                                if rel.RelatingObject.is_a("IfcElement")}))
        rows.append(Row(identity(entity.GlobalId), entity.Name or "", code, tid,
                        facts(entity), types.get(tid, []), typ.Name or "" if typ else "",
                        parents=parents))
    clock.mark("ifcRows")
    clock.seconds["ifcRows"] -= pset_seconds
    clock.seconds["psetRead"] = pset_seconds
    if timings is not None:
        timings.update(clock.seconds)
    return rows


def read_usd(path):
    """Read quarantined camera facts from a published core stage, without IFC."""
    stage = Usd.Stage.Open(str(path))
    if not stage or stage.GetCompositionErrors():
        raise ValueError("Source USD stage must compose successfully")
    rows = []
    for prim in stage.Traverse():
        if not prim.HasAPI("AecoElementAPI"):
            continue
        cat = _catalog(prim)
        rows.append(Row(identity(prim.GetAttribute("aeco:id").Get()),
                        prim.GetDisplayName() or prim.GetName(),
                        prim.GetAttribute("aeco:class:ifc:code").Get() or "",
                        str(cat.GetPath()) if cat else "",
                        _quarantine(prim), _quarantine(cat),
                        cat.GetName() if cat else ""))
    return rows


def read_cobie(path):
    """Read Component/Type/Attribute using headers, independent of column order.

    Attribute Category optionally preserves the original Pset. Unit applies
    to source lengths only. ExtIdentifier, IfcGUID and IfcGlobalId are accepted.
    """
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        def sheet(name):
            if name not in book:
                raise ValueError("COBie requires " + name + " sheet")
            cells = book[name].iter_rows(values_only=True)
            headers = [_norm(v or "") for v in next(cells)]
            return [dict(zip(headers, row)) for row in cells if any(v is not None for v in row)]
        components, types, attributes = sheet("Component"), sheet("Type"), sheet("Attribute")
    finally:
        book.close()
    facts = defaultdict(list)
    for a in attributes:
        if not a.get("name"):
            continue
        units = _norm(a.get("unit") or "mm")
        scale = {"m": 1., "metre": 1., "meter": 1., "meters": 1., "metres": 1.,
                 "cm": .01, "mm": .001, "millimeters": .001, "millimetres": .001,
                 "ft": .3048, "feet": .3048, "in": .0254}.get(units, .001)
        category = a.get("category")
        pset = str(category) if _present(category) else "COBie"
        if _norm(a["name"]) in {_norm(n) for n in ("IsOutdoors", "VideoResolutionWidth", "VideoResolutionHeight",
                                                     "TiltHorizontal", "PanHorizontal", "Zoom", "PanTiltZoomPreset")}:
            pset = CAMERA_PSET
        facts[(str(a.get("sheetname") or "").lower(), str(a.get("rowname") or ""))].append(
            Fact(pset, str(a["name"]), a.get("value"), scale))
    by_type = {str(t.get("name")): t for t in types}
    component_names = Counter(str(c.get("name") or "") for c in components)
    rows = []
    for c in components:
        name, type_name = str(c.get("name") or ""), str(c.get("typename") or "")
        own = facts[("component", name)]
        typ = by_type.get(type_name, {})
        candidates = [c.get(k) for k in ("extidentifier", "ifcguid", "ifcglobalid")]
        candidates += [f.value for f in own if _norm(f.name) in ("ifcguid", "ifcglobalid")]
        uid = next((identity(v) for v in candidates if _present(v) and len(identity(v)) == 36), "")
        code = c.get("extobject") or typ.get("extobject") or "IfcBuildingElementProxy"
        rows.append(Row(uid, name, str(code), identity(typ.get("extidentifier")) or type_name,
                        own, facts[("type", type_name)], type_name, component_names[name] > 1))
    return rows


def _helper_optics(facts):
    """The graphics helper dialect; these observations never become drivers."""
    aliases = {"focallengthminimum": "FOV Focal Length Minimum",
               "focallengthmaximum": "FOV Focal Length Maximum",
               "horizontalres": "FOV Horizontal Resolution",
               "verticalres": "FOV Vertical Resolution"}
    values = _optics(Values([Fact(f.pset, aliases.get(_norm(f.name), f.name),
                                 f.value, f.length_scale) for f in facts]))
    keys = ("focalRange", "hfovRange", "vfovRange", "pixels")
    if all(key in values and all(v > 0 for v in values[key]) for key in keys):
        return {key: values[key] for key in keys}
    return {}


def _is_sub(row, facts):
    # A real CAMERA remains a referent even when nested or named like a picture.
    if row.code == "IfcAudioVisualAppliance.CAMERA":
        return False
    if Values(facts).pick("SuperComponent") is not None:
        return True
    names = [_norm(row.name), _norm(row.type_name)] + [
        _norm(str(f.value)) for f in facts if isinstance(f.value, str)
        and _norm(f.name) in ("name", "family", "familyname", "typename", "objecttype")]
    names += [_norm(f.pset) for f in facts]
    if any("fovaxis" in n or "axis2dsymbol" in n for n in names):
        return True
    # Revit can replace the family name with a generic proxy label. Require
    # helper-specific graphics/formula evidence, not a camera's symbol choice.
    return (row.code.split(".")[0] == "IfcBuildingElementProxy"
            and any(_norm(f.name).startswith("rglength") for f in facts)
            and bool(_helper_optics(facts)))


def is_camera(code, facts, *, classes=None):
    """Proxy classes require optical evidence, never just a product-like name."""
    if code.split(".")[0] == "IfcAudioVisualAppliance":
        return code == "IfcAudioVisualAppliance.CAMERA"
    classes = registry("camera_classes") if classes is None else classes
    if code not in classes and code.split(".")[0] not in classes:
        return False
    if code.split(".")[0] == "IfcBuildingElementProxy":
        return any(_norm(f.pset) == _norm(exchange.PSET) and _norm(f.name) == "sensors" and _present(f.value) for f in facts) or any(_norm(f.name) in {_norm(n) for n in (
            "FOV Desired Focal Length", "FOV Actual Focal Length", "FOV Focal Length Minimum",
            "FOV Pan", "FOV Camera Rotation", "FOV 1 Pan")} and _present(f.value) for f in facts)
    return True


def _optics(values):
    result = contract.optics(values)
    for key, make in (("focalRange", Gf.Vec2d), ("hfovRange", Gf.Vec2d),
                      ("vfovRange", Gf.Vec2d), ("pixels", Gf.Vec2i), ("offset", Gf.Vec3d)):
        if key in result:
            result[key] = make(*result[key])
    return result


def _device(prim, values):
    outdoors = values.pick("IsOutdoors", pset=CAMERA_PSET, convert=_bool)
    if outdoors is not None:
        _set(prim, "aeco:cctvType:outdoor", outdoors)
    before = set(values.used)
    placement = values.pick("Placement ID", convert=lambda v: int(_number(v)))
    mount = MOUNTS.get(placement)
    if placement is not None and mount is None:
        values.used = before
    if placement is None:
        text = values.pick("Mounting Type")
        if text:
            mount = next((n for n in ("corner", "pole", "wall", "pendant", "recessed", "ceiling",
                                      "parapet", "desk") if n in str(text).lower()), "other")
    _set(prim, "aeco:cctv:mount", mount)
    # Unknown scenarios remain quarantined instead of being silently lost.
    candidate = Values(values.facts)
    scenario = candidate.pick("Scenario")
    if scenario is not None and str(scenario).lower() in {v.lower() for v in registry("scenarios")}:
        _set(prim, "aeco:cctv:scenario", str(scenario).lower())
        values.used.update(candidate.used)
    code = values.pick("Classification.Uniclass.Pr.Number")
    if code:
        prim.ApplyAPI("AecoClassificationAPI", "uniclass")
        _set(prim, "aeco:class:uniclass:code", str(code))


def _collapse_types(rows, index, *, camera_classes=None):
    """Share exact duplicate source catalogs, preserving all occurrence facts.

    Equal optics alone cannot justify discarding a different manufacturer,
    mount or unknown type property. Compare the complete type record and the
    composed core catalog too; catalogs with children/arcs are left intact.
    """
    groups = defaultdict(list)
    seen = set()
    layouts = defaultdict(set)
    for row in rows:
        facts = row.facts + row.type_facts
        numbers = {int(m.group(1)) for f in facts
                   if (m := re.match(r"fov(\d+)(?:pan|tilt|desiredfocallength|actualfocallength)$", _norm(f.name)))}
        presets = any(re.fullmatch(r"preset\d+", _norm(f.name)) for f in facts)
        table = Values(facts).pick("PanTiltZoomPreset", pset=CAMERA_PSET)
        count = 1 if presets or table is not None else max(numbers, default=1)
        # Contract payloads can specify head topology independently of the
        # legacy names. Leave their source catalogs intact in this pass.
        contract = any(_norm(f.pset) == _norm(exchange.PSET) for f in facts)
        layouts[row.type_id].add((count, contract))
    for row in rows:
        prim = index.get(row.uid)
        if (not prim or not row.type_id or not row.type_name or not row.type_facts
                or _is_sub(row, row.facts + row.type_facts)
                or not is_camera(row.code, row.facts + row.type_facts, classes=camera_classes)):
            continue
        cat = _catalog(prim)
        if cat is None or (row.type_id, cat.GetPath()) in seen:
            continue
        seen.add((row.type_id, cat.GetPath()))
        if len(layouts[row.type_id]) != 1 or next(iter(layouts[row.type_id]))[1]:
            continue
        if (cat.GetAllChildren() or cat.GetInherits().GetAllDirectInherits()
                or cat.HasAuthoredReferences() or cat.HasAuthoredPayloads()
                or cat.GetRelationships()
                or any(a.GetNumTimeSamples() for a in cat.GetAttributes())):
            continue
        source = sorted((f.pset, f.name, json.dumps(f.value, sort_keys=True), f.length_scale)
                        for f in row.type_facts)
        properties = [(a.GetName(), str(a.GetTypeName()), str(a.Get())) for a in cat.GetAttributes()]
        # Include resolved optics so conflicting aliases cannot be reordered
        # into equality by the source-record sort.
        optics = {k: str(v) for k, v in _optics(Values(row.type_facts)).items()}
        key = json.dumps([row.code, row.type_name, source, properties, optics,
                          cat.GetAppliedSchemas(), sorted(layouts[row.type_id])], sort_keys=True)
        groups[key].append((row.type_id, cat.GetPath()))
    replacements, report = {}, {}
    for group in groups.values():
        if len(group) < 2:
            continue
        canonical_id, canonical_path = min(group)
        for source_id, source_path in sorted(group):
            replacements[source_path] = canonical_path
            report[source_id] = dict(typeId=canonical_id, catalog=str(canonical_path),
                                     sourceCatalog=str(source_path))
    for row in rows:
        prim = index.get(row.uid)
        if not prim:
            continue
        cat = _catalog(prim)
        if cat is not None and cat.GetPath() in replacements:
            paths = prim.GetInherits().GetAllDirectInherits()
            prim.GetInherits().SetInherits(list(dict.fromkeys(replacements.get(p, p) for p in paths)))
    return report


def _reverse_subinstances(mapping):
    reverse = {}
    for owner, children in mapping.items():
        if not isinstance(children, list):
            raise ValueError("subinstances values must be lists of child GUIDs")
        for child in children:
            key = identity(child)
            if key in reverse:
                raise ValueError("Sub-instance has multiple owners")
            reverse[key] = identity(owner)
    return reverse


def _owned_camera(uid, rows, camera_ids, visiting=None):
    """Follow both IFC decomposition relations, including intermediate helpers."""
    visiting = set() if visiting is None else visiting
    if uid in visiting:
        raise ValueError("Sub-instance ownership cycle: " + uid)
    if uid in camera_ids:
        return {uid}
    row = rows.get(uid)
    owners = set()
    for parent in row.parents if row else ():
        owners.update(_owned_camera(parent, rows, camera_ids, visiting | {uid}))
    if len(owners) > 1:
        raise ValueError("Sub-instance has multiple IFC camera owners: " + uid)
    return owners


def import_cctv(core_stage, source_file, output="kind.usda", *, subinstances=None, timings=None):
    """Return exactly STATS counters and write a new layer sublayering core.

    Optional subinstances maps camera GUIDs to child GUID lists. IFC ownership
    then SuperComponent precede same-level proximity (0.5 m, or 1 m with a
    matching complete helper optics observation). No geometry is authored.
    Duplicate catalog mappings are JSON in layer custom data, leaving STATS
    unchanged for existing callers.
    Optional timings receives exclusive stage seconds and total wall time.
    """
    clock = _Timings()
    register_plugins()
    core_stage, source_file, output = map(lambda p: Path(p).resolve(), (core_stage, source_file, output))
    if output.exists() or output in (core_stage, source_file):
        raise ValueError("Output must be a new file distinct from the inputs")
    clock.mark("setup")
    source_times = {}
    rows = (read_cobie(source_file) if source_file.suffix.lower() == ".xlsx"
            else read_usd(source_file) if source_file.suffix.lower() in (".usd", ".usda", ".usdc")
            else read_ifc(source_file, timings=source_times))
    clock.mark("sourceRead")
    clock.seconds["sourceRead"] -= sum(source_times.values())
    clock.seconds.update(source_times)
    reverse = _reverse_subinstances(subinstances or {})
    base = Usd.Stage.Open(str(core_stage))
    if not base or base.GetCompositionErrors():
        raise ValueError("Core stage must compose successfully")
    layer = Sdf.Layer.CreateAnonymous("kind.usda")
    layer.subLayerPaths = [str(core_stage)]
    stage = Usd.Stage.Open(layer)
    for key in ("defaultPrim", "upAxis", "metersPerUnit", "fallbackPrimTypes"):
        if base.HasAuthoredMetadata(key):
            stage.SetMetadata(key, base.GetMetadata(key))
    layer.customLayerData = {"aeco:library": "usdAecoCctv", "aeco:version": __version__}
    clock.mark("stageOpen")
    index = {}
    for prim in stage.Traverse():
        uid = prim.GetAttribute("aeco:id").Get()
        if uid:
            uid = identity(uid)
            if uid in index:
                raise ValueError("Core stage has duplicate aeco:id")
            index[uid] = prim
    clock.mark("primTraversal")
    stats = Counter(dict.fromkeys(STATS, 0))
    camera_classes = registry("camera_classes")
    cameras, subs, done_types = {}, [], set()
    # Snapshot before blocking catalog opinions; occurrences may inherit them.
    requested = {row.uid for row in rows}
    snapshots = {uid: _quarantine(p) for uid, p in index.items() if uid in requested}
    clock.mark("quarantineRead")
    type_mapping = _collapse_types(rows, index, camera_classes=camera_classes)
    if type_mapping:
        layer.customLayerData = {**layer.customLayerData,
                                 "aeco:cctv:typeMapping": json.dumps(type_mapping, sort_keys=True)}
    clock.mark("typeMatching")
    for row in rows:
        prim = index.get(row.uid)
        facts = row.facts + row.type_facts + snapshots.get(row.uid, [])
        if _is_sub(row, facts) or (row.uid in reverse and row.code != "IfcAudioVisualAppliance.CAMERA"):
            subs.append((row, prim, facts))
            continue
        code = row.code if row.code.startswith("IfcAudioVisualAppliance.") else (prim.GetAttribute("aeco:class:ifc:code").Get() if prim else row.code)
        if not is_camera(code or row.code, facts, classes=camera_classes):
            continue
        if row.ambiguous:
            raise ValueError("COBie camera Component.Name must be unique for Attribute matching")
        if prim is None:
            stats["unmatched"] += 1
            continue
        if not prim.HasAPI("AecoElementAPI"):
            raise ValueError("Camera must be a core element")
        if row.uid in cameras:
            raise ValueError("Source has duplicate camera identity")
        cameras[row.uid] = prim
        prim.ApplyAPI("AecoCctvCameraAPI")
        stats["cameras"] += 1
        cat = _catalog(prim)
        if cat is None:
            root = stage.GetDefaultPrim().GetPath() if stage.GetDefaultPrim() else Sdf.Path.absoluteRootPath
            catpath = root.AppendChild("_TypeCatalog").AppendChild(
                "Camera_" + re.sub(r"[^A-Za-z0-9_]", "_", row.type_id or row.uid))
            stage.CreateClassPrim(catpath.GetParentPath())
            cat = stage.GetPrimAtPath(catpath) or stage.CreateClassPrim(catpath)
            cat.ApplyAPI("AecoTypeAPI")
            prim.GetInherits().AddInherit(catpath)
        cat.ApplyAPI("AecoCctvCameraTypeAPI")
        # CameraAPI on the class gives the housing's mount a typed home.
        cat.ApplyAPI("AecoCctvCameraAPI")
        type_values = Values(row.type_facts + _quarantine(cat))
        type_a = exchange.decode(type_values, row.type_name or cat.GetName())
        own_values = Values(row.facts)
        own_a = exchange.decode(own_values, row.name)
        values = Values(facts)
        numbers = sorted({int(m.group(1)) for f in facts
                          if (m := re.match(r"fov(\d+)(?:pan|tilt|desiredfocallength|actualfocallength)$", _norm(f.name)))})
        flags = sorted({int(m.group(1)) for f in facts if (m := re.fullmatch(r"preset(\d+)", _norm(f.name)))})
        table = values.pick("PanTiltZoomPreset", pset=CAMERA_PSET)
        presets = _table_presets(table) if table is not None else {}
        for n in flags:
            flag = values.pick("Preset %d" % n, convert=_bool)
            if flag:
                presets.setdefault("Preset_%d" % n, _pose(values, n))
        count = 1 if flags or table is not None else max(numbers, default=1)
        type_optics = _optics(type_values)
        own_optics = _optics(values)
        if not row.type_facts and not type_optics:
            type_optics = dict(own_optics)
        if presets and values.pick("PTZ", convert=_bool) is None:
            own_optics["motorised"] = True
        common = _pose(values)
        _device(cat, type_values)
        _device(prim, values)
        type_heads, own_heads = type_a.get("Sensors", []), own_a.get("Sensors", [])
        if "Sensors" in type_a or "Sensors" in own_a:
            count = len(type_heads or own_heads)
            if count == 0 or (type_heads and own_heads and len(type_heads) != len(own_heads)):
                raise ValueError(row.name + ": Sensor_0 missing drivers or mismatched sensor lists")
        for target, payload in ((cat, type_a), (prim, own_a)):
            for key, value in {**payload.get("Type", {}), **payload.get("Drivers", {})}.items():
                if key in exchange.HOUSING:
                    if key == "aeco:cctvType:outdoor":
                        mirror = values.pick("IsOutdoors", pset=CAMERA_PSET, convert=_bool)
                        if mirror is not None:
                            exchange.conflict(row.name, key, value, mirror)
                    _set(target, key, value)
        for i in range(count):
            ta = type_heads[i] if type_heads else {}
            oa = own_heads[i] if own_heads else {}
            td, od = exchange.drivers(ta), exchange.drivers(oa)
            path = cat.GetPath().AppendChild("Sensor_%d" % i)
            ts = stage.GetPrimAtPath(path)
            if not ts:
                ts = stage.DefinePrim(path, "Camera")
                ts.ApplyAPI("AecoCctvSensorAPI")
            for key, value in {**type_optics, **td}.items():
                _set(ts, SENSOR + key, value)
            sensor = stage.GetPrimAtPath(prim.GetPath().AppendChild("Sensor_%d" % i))
            for key, value in {**own_optics, **td, **od}.items():
                _set(sensor, SENSOR + key, value)
            pose = dict(common)
            if numbers and not flags and table is None:
                pose.update(_pose(values, i + 1))
            pose.update(td)
            pose.update(od)
            mirror_values = Values([f for f in facts if _norm(f.pset) == _norm(CAMERA_PSET)])
            mirrors = {**_optics(mirror_values), **_pose(mirror_values)} if i == 0 else {}
            for key, value in {**td, **od}.items():
                if key in mirrors:
                    exchange.conflict(str(sensor.GetPath()), SENSOR + key, value, mirrors[key])
            for key, value in pose.items():
                _set(sensor, SENSOR + key, value)
            head_presets = _table_presets(table, i) if table is not None else dict(presets)
            for entry in (ta, oa):
                if "presets" in entry:
                    head_presets = entry["presets"]
                if "tour" in entry:
                    _set(sensor, SENSOR + "tour", entry["tour"])
            if not isinstance(head_presets, dict):
                raise ValueError(str(sensor.GetPath()) + ": presets must be an object")
            for old in list(sensor.GetAppliedSchemas()):
                if old.startswith("AecoCctvPresetAPI:"):
                    sensor.RemoveAPI("AecoCctvPresetAPI", old.split(":", 1)[1])
            for name, drivers in head_presets.items():
                if not Sdf.Path.IsValidIdentifier(name):
                    raise ValueError(str(sensor.GetPath()) + ": invalid preset name " + name)
                sensor.ApplyAPI("AecoCctvPresetAPI", name)
                for key, value in drivers.items():
                    if key in ("pan", "tilt", "focalLength", "dwell", "home"):
                        _set(sensor, "aeco:cctvPreset:%s:%s" % (name, key), value)
                stats["presets"] += 1
            from .validators import missing_optics
            missing = missing_optics(sensor)
            if missing:
                raise ValueError("cctvMissingOptics: %s missing drivers: %s" % (sensor.GetPath(), ", ".join(missing)))
            stats["sensors"] += 1
        if cat.GetPath() not in done_types:
            stats["types"] += 1
            done_types.add(cat.GetPath())
        _block(prim, values.used | type_values.used | own_values.used, stats)
        _block(cat, type_values.used, stats)
    clock.mark("cameraAuthoring")
    camera_ids = {row.uid for row in rows if is_camera(row.code, row.facts + row.type_facts, classes=camera_classes)
                  and not _is_sub(row, row.facts + row.type_facts) and row.uid not in reverse}
    by_id = {row.uid: row for row in rows}
    for row, _, _ in subs:
        owners = _owned_camera(row.uid, by_id, camera_ids)
        if row.uid not in reverse and owners:
            reverse[row.uid] = next(iter(owners))
    _fold(stage, cameras, subs, reverse, stats)
    clock.mark("helperMatching")
    output.parent.mkdir(parents=True, exist_ok=True)
    exported = Sdf.Layer.CreateAnonymous("kind-export.usda")
    exported.TransferContent(layer)
    exported.subLayerPaths = [os.path.relpath(core_stage, output.parent)]
    publish(exported, output)
    clock.mark("writing")
    clock.finish(timings)
    return dict(stats)


def _level(prim):
    while prim and not prim.IsPseudoRoot():
        if prim.GetTypeName() == "AecoLevel":
            return prim.GetPath()
        prim = prim.GetParent()
    return None


class _CameraIndex:
    """Metre-sized cells per level, with resolved optics cached per sensor."""
    KEYS = ("focalRange", "hfovRange", "vfovRange", "pixels")
    TOLERANCE = 1e-6

    def __init__(self, cameras, cache, metres):
        self.origins, self.sensors, self.heads = {}, {}, {}
        self.cells = defaultdict(list)
        optics = []
        for uid, prim in cameras.items():
            point = cache.GetLocalToWorldTransform(prim).ExtractTranslation() * metres
            self.origins[uid] = point
            self.sensors[uid] = sensors_of(prim)
            heads = []
            for sensor in self.sensors[uid]:
                heads.append(len(optics))
                optics.append([v for key in self.KEYS
                               for v in sensor.GetAttribute(SENSOR + key).Get()])
            self.heads[uid] = heads
            level = _level(prim)
            if level is not None and all(math.isfinite(v) for v in point):
                self.cells[(level, *(math.floor(v) for v in point))].append(uid)
        self.optics = np.asarray(optics, dtype=float).reshape((-1, 8))

    def near(self, point, level, observations, limit):
        if not all(math.isfinite(v) for v in point):
            return []
        # A second camera just beyond the radius can still tie the nearest
        # eligible one. Include the tie halo before applying the final bound.
        radius = limit + self.TOLERANCE
        axes = [range(math.floor(v - radius), math.floor(v + radius) + 1) for v in point]
        candidates = [uid for cell in product(*axes)
                      for uid in self.cells.get((level, *cell), ())]
        if observations and candidates:
            owners = [uid for uid in candidates for _ in self.heads[uid]]
            heads = [head for uid in candidates for head in self.heads[uid]]
            observed = np.asarray([v for key in self.KEYS for v in observations[key]])
            matched = np.all(np.abs(self.optics[heads] - observed) <= self.TOLERANCE, axis=1)
            candidates = set(uid for uid, match in zip(owners, matched) if match)
        # Keep Gf's distance arithmetic and UUID ordering for byte-identical
        # distanceMetres metadata, including ties and negative coordinates.
        return sorted((distance, uid) for uid in candidates
                      if (distance := (self.origins[uid] - point).GetLength()) <= radius)


def _fold(stage, cameras, subs, reverse, stats):
    cache = UsdGeom.XformCache()
    metres = UsdGeom.GetStageMetersPerUnit(stage)
    index = _CameraIndex(cameras, cache, metres)
    groups = defaultdict(list)
    report = {}
    for row, prim, facts in subs:
        owner = reverse.get(row.uid)
        level = _level(prim)
        method, distance = "ownership", None
        if owner is None:
            handle = Values(facts).pick("SuperComponent")
            owner = identity(handle) if handle else None
            method = "SuperComponent"
        # An explicit owner that is missing from core must stay unmatched;
        # proximity must not silently assign its picture to a different device.
        if owner is None and prim and level is not None:
            point = cache.GetLocalToWorldTransform(prim).ExtractTranslation() * metres
            method, limit = "nearest", .5
            observations = _helper_optics(facts)
            if observations:
                # The exported picture's origin can be below the housing.
                # Expand only with complete matching optical evidence, and
                # retain a metre bound and the unique-nearest requirement.
                method, limit = "nearest-optics", 1.
            near = index.near(point, level, observations, limit)
            if near and near[0][0] <= limit and (len(near) == 1 or near[1][0] - near[0][0] > 1e-6):
                distance, owner = near[0]
        report[row.uid] = dict(owner=owner or "", method=method,
                               matched=bool(owner in cameras and prim is not None))
        if distance is not None:
            report[row.uid]["distanceMetres"] = distance
        if owner not in cameras or prim is None:
            report[row.uid]["reason"] = ("missing core helper" if prim is None else
                "owner is not an imported camera" if owner else
                "helper has no level" if level is None else
                "no unique camera within the matching distance")
            stats["unmatched"] += 1
            continue
        groups[owner].append((row, prim))
    has_sync = Usd.SchemaRegistry().FindAppliedAPIPrimDefinition("AecoHostBindingAPI") is not None
    for owner, children in groups.items():
        sensors = index.sensors[owner]
        for i, (row, prim) in enumerate(sorted(children, key=lambda pair: pair[0].uid)):
            sensor = sensors[min(i, len(sensors) - 1)]
            kind = "symbol" if "axis2dsymbol" in _norm(row.name) else "fov"
            ref = owner + ":" + kind + ":" + row.uid
            # The sync contract has one scalar ref. Keep the first there and
            # retain all extra FOV/preset/symbol handles in the string array.
            if has_sync and not sensor.HasAPI("AecoHostBindingAPI", "revit"):
                sensor.ApplyAPI("AecoHostBindingAPI", "revit")
                sensor.GetAttribute("aeco:host:revit:ref").Set(ref)
            else:
                attr = sensor.CreateAttribute("aeco:props:cctv:subInstances", Sdf.ValueTypeNames.StringArray)
                attr.Set(sorted(set(list(attr.Get() or []) + [ref])))
            prim.SetActive(False)
            stats["subInstancesFolded"] += 1
    if report:
        layer = stage.GetRootLayer()
        layer.customLayerData = {**layer.customLayerData,
                                 "aeco:cctv:foldMapping": json.dumps(report, sort_keys=True)}


def main(argv=None):
    parser = argparse.ArgumentParser(prog="aeco-cctv-import", description=__doc__)
    parser.add_argument("core_stage")
    parser.add_argument("source", nargs="?", help="IFC, COBie or published USD; defaults to core_stage")
    parser.add_argument("-o", "--out", default="kind.usda")
    parser.add_argument("--subinstances", type=Path, help="optional JSON camera GUID -> child GUID list")
    args = parser.parse_args(argv)
    try:
        mapping = json.loads(args.subinstances.read_text()) if args.subinstances else None
        timings = {}
        stats = import_cctv(args.core_stage, args.source or args.core_stage, args.out, subinstances=mapping, timings=timings)
    except (ValueError, RuntimeError, OSError) as exc:
        parser.exit(1, "aeco-cctv-import: " + str(exc) + "\n")
    print(json.dumps({**stats, "timingsSeconds": timings}, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
